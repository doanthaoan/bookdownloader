"""
Translation service for raw TXT novels.
Translates Chinese raw text to Vietnamese via the dichtienghoa.com API,
with a Selenium web-method fallback if the API is not callable.
"""

import os
import re
import time
import random
import threading
import requests
from pathlib import Path
from urllib.parse import urlparse

from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from colorama import Fore, init

from app.config import TRUYENWIKI, get_user_agent
from app.database import get_database

init(autoreset=True)

# Track active translation jobs for progress/cancel
_active_translations = {}
_lock = threading.Lock()

# Chapter title pattern used when parsing raw TXT files.
# Matches lines like "12. 第1章 xxx" or "第1章 xxx"
CHAPTER_SPLIT_RE = re.compile(r"(?=\b\d+\.\s第\d+[章话])|(?=^第\s*\d+\s*[章话])", re.MULTILINE)

# Split text into translation chunks at paragraph boundaries (avoids very long single requests).
MAX_CHUNK_CHARS = 1200


class CloudflareBlockedError(RuntimeError):
    """Raised when a translation request cannot pass the Cloudflare check after
    all cookie-refresh attempts. The run should STOP at the current chapter
    (marking it failed) so the user can refresh the cookie and click Continue —
    it must NOT move on to the next chapter and fail the whole batch."""


class TranslateTask:
    def __init__(self, book_id: int):
        self.book_id = book_id
        self.cancelled = False
        self.progress = {
            "active": True,
            "success_count": 0,
            "fail_count": 0,
            "total": 0,
            "current_index": 0,
            "current_title": "",
            "cancelled": False,
            "method": "api",
            "message": "",
        }

    def set(self, **kwargs):
        self.progress.update(kwargs)


def register_task(task: TranslateTask):
    with _lock:
        _active_translations[task.book_id] = task


def unregister_task(book_id: int):
    with _lock:
        _active_translations.pop(book_id, None)


def cancel_translation(book_id: int) -> bool:
    with _lock:
        task = _active_translations.get(book_id)
    if not task:
        return False
    task.cancelled = True
    task.progress["cancelled"] = True
    return True


def get_translation_progress(book_id: int):
    with _lock:
        task = _active_translations.get(book_id)
    return task.progress if task else None


def translate_text(text: str, method: str = "api") -> str:
    """Translate a block of text. method: 'api' or 'web'."""
    return Translator().translate(text, method=method)


def _read_text_file(path: str) -> str:
    """Read a text file, trying multiple encodings. Handles UTF-8 BOM,
    UTF-16 (LE/BE), and common Chinese encodings (GBK)."""
    raw = None
    with open(path, "rb") as f:
        raw = f.read()
    # Strip BOMs if present
    if raw.startswith(b"\xef\xbb\xbf"):
        raw = raw[3:]
        text = raw.decode("utf-8")
        return text.replace("\r\n", "\n").replace("\r", "\n")
    for enc in ("utf-8", "utf-16", "gbk", "gb18030", "big5", "latin-1"):
        try:
            text = raw.decode(enc)
            return text.replace("\r\n", "\n").replace("\r", "\n")
        except (UnicodeDecodeError, LookupError):
            continue
    # Last resort: latin-1 never fails
    return raw.decode("latin-1").replace("\r\n", "\n").replace("\r", "\n")


def parse_raw_file(path: str) -> dict:
    """Parse a raw TXT novel into book info + chapters.
    Returns {'raw_title', 'author', 'summary', 'chapters': [{order, title, content}]}.
    'summary' is the intro/meta text before the first chapter (minus title/author
    lines) — used as the book summary, translated on save.
    """
    content = _read_text_file(path)

    parts = [p.strip() for p in CHAPTER_SPLIT_RE.split(content) if p.strip()]
    if not parts:
        return {"raw_title": os.path.basename(path), "author": "", "summary": "", "chapters": []}

    meta = parts[0]
    lines = [l.strip() for l in meta.split("\n") if l.strip()]

    raw_title = lines[0] if lines else os.path.basename(path)
    # Try to detect an author line from common patterns
    author = ""
    summary_lines = []
    for l in lines[1:]:
        m = re.match(r"^\s*(?:Tác giả|作者|著者)\s*[:：]?\s*(.+)$", l)
        if m and m.group(1) and not author:
            author = m.group(1).strip()
            continue
        summary_lines.append(l)

    chapters = []
    for i, part in enumerate(parts[1:], 1):
        sub = part.split("\n", 1)
        title = sub[0].strip()
        content = sub[1].strip() if len(sub) > 1 else ""
        chapters.append({"order": i, "title": title, "content": content})

    return {
        "raw_title": raw_title,
        "author": author,
        "summary": "\n".join(summary_lines).strip(),
        "chapters": chapters,
    }


def apply_corrections(text: str, corrections: list) -> str:
    """Apply per-book find/replace corrections to translated text.

    Matching is case-insensitive so a name at the start of a paragraph or
    right after an opening quote (which Vietnamese capitalizes) still matches
    the lower-cased find_text the user entered. The replacement text is used
    verbatim (e.g. 'lộ khi' -> 'Lộ Thời').
    """
    for c in corrections:
        if not c.get("enabled", True):
            continue
        find_text = c.get("find_text") or ""
        replace_text = c.get("replace_text") or ""
        if find_text:
            text = re.sub(re.escape(find_text), lambda m: replace_text, text, flags=re.IGNORECASE)
    return text


class Translator:
    """Translates text using dichtienghoa API or web method."""

    def __init__(self):
        self.db = get_database()
        self.api_endpoint = self.db.get_setting("translate_api_endpoint") or "https://dichtienghoa.com/transtext"
        self.site = self.db.get_setting("translate_site") or "https://dichtienghoa.com"
        self.target_lang = self.db.get_setting("translate_target_lang") or "vi"
        self.click_delay = float(self.db.get_setting("translate_click_delay") or 3)
        self.retry_delay = float(self.db.get_setting("translate_retry_delay") or 10)
        self.max_retries = int(self.db.get_setting("translate_max_retries") or 3)
        self.browser_ua = (self.db.get_setting("translate_browser_ua") or "").strip() or get_user_agent()
        self.impersonate = (self.db.get_setting("translate_impersonate") or "chrome").strip() or "chrome"
        self.cookies = self._load_cookies()
        self.driver = None
        # Auto cookie-refresh when Cloudflare blocks a request. A real Chrome
        # (Selenium) is opened, the challenge is passed (possibly auto), and the
        # fresh cf_clearance is harvested and saved so the process does not stop.
        self.cookie_wait = float(self.db.get_setting("translate_cookie_wait") or 60)
        self.cookie_refresh_max = int(self.db.get_setting("translate_cookie_refresh_max") or 1)
        self._cookie_refresh_count = 0

    def _load_cookies(self) -> dict:
        raw = self.db.get_setting("translate_cookies") or "{}"
        try:
            import json
            return json.loads(raw) or {}
        except Exception:
            return {}

    def _save_cookies(self, cookies: dict):
        """Persist a fresh cookie set to settings AND update this instance."""
        import json
        self.cookies = cookies
        self.db.update_setting("translate_cookies", json.dumps(cookies))

    # --- API method ---

    def api_available(self) -> bool:
        """Quickly probe whether the API endpoint responds to a tiny request."""
        try:
            r = self._api_call("你好", timeout=15)
            return bool(r and r.strip())
        except Exception as e:
            print(f"{Fore.YELLOW}Translation API probe failed: {e}")
            return False

    def _request_api(self, data: dict, headers: dict, timeout: int):
        """Issue the API POST request using curl_cffi (with fallback to requests)."""
        try:
            from curl_cffi import requests as cffi_requests
            return cffi_requests.post(
                self.api_endpoint, data=data, headers=headers,
                cookies=self.cookies or None, impersonate=self.impersonate,
                timeout=timeout,
            )
        except ImportError:
            return requests.post(self.api_endpoint, data=data, headers=headers,
                                 cookies=self.cookies or None, timeout=timeout)

    def _api_call(self, text: str, timeout: int = 60) -> str:
        """Send form-encoded request to the translation API.
        Expected response: {"data": "...", "err": 0}

        Uses curl_cffi with Chrome TLS impersonation + the real browser UA so a
        stored cf_clearance cookie (bound to that fingerprint) can pass Cloudflare.
        Falls back to plain requests if curl_cffi is unavailable.

        If Cloudflare blocks (403), the cookie is auto-refreshed from a real
        Chrome (up to cookie_refresh_max times) and the request is retried.
        """
        headers = {
            "User-Agent": self.browser_ua,
            "Accept": "application/json, text/javascript, */*; q=0.01",
            "Content-Type": "application/x-www-form-urlencoded; charset=UTF-8",
            "X-Requested-With": "XMLHttpRequest",
            "Referer": self.site + "/",
        }
        data = {"t": text, "tt": self.target_lang}

        attempts = self.cookie_refresh_max + 1
        for attempt in range(1, attempts + 1):
            r = self._request_api(data, headers, timeout)
            if r.status_code == 403 and ("Just a moment" in (r.text or "") or "cf_clearance" in (r.text or "")):
                print(f"{Fore.YELLOW}[api] Cloudflare block (attempt {attempt}/{attempts}). Refreshing cookie...")
                if not self._maybe_refresh_cookies():
                    break
                continue
            r.raise_for_status()
            result = r.json()
            if result.get("err") == 0:
                return str(result.get("data") or "").strip()
            raise RuntimeError(f"Translation API error: {result}")
        raise CloudflareBlockedError(
            f"API translation blocked by Cloudflare after {attempts} attempt(s). "
            f"Auto cookie refresh limit ({self.cookie_refresh_max}) reached. "
            f"Update translate_cookies under Settings (or click 'Continue' on the book)."
        )

    # --- Web method (Selenium) ---

    def _ensure_driver(self):
        if self.driver is not None:
            return
        from selenium import webdriver
        from selenium.webdriver.chrome.service import Service
        from selenium.webdriver.chrome.options import Options
        from webdriver_manager.chrome import ChromeDriverManager

        options = Options()
        options.add_argument("--no-sandbox")
        options.add_argument("--disable-dev-shm-usage")
        options.add_argument("--disable-blink-features=AutomationControlled")
        # NOTE: do NOT force a custom user-agent here. cf_clearance cookies are
        # bound to the exact UA + TLS fingerprint that solved the Cloudflare
        # challenge. Chrome's real UA must match the user's browser.
        self.driver = webdriver.Chrome(
            service=Service(ChromeDriverManager().install()),
            options=options
        )

    def _inject_cookies(self, driver):
        """Inject stored cookies (e.g. cf_clearance) into the current page."""
        if not self.cookies:
            return
        domain = urlparse(self.site).hostname
        for name, value in self.cookies.items():
            try:
                driver.add_cookie({"name": name, "value": value, "domain": domain})
            except Exception as e:
                print(f"{Fore.YELLOW}Could not add translate cookie {name}: {e}")

    def _maybe_refresh_cookies(self) -> bool:
        """Open a real Chrome, let Cloudflare issue a fresh cf_clearance, save it.
        Guarded by cookie_refresh_max so we never loop forever (anti-blocking).
        """
        if self._cookie_refresh_count >= self.cookie_refresh_max:
            print(f"{Fore.RED}[cookie] Refresh limit reached ({self.cookie_refresh_max}).")
            return False
        self._cookie_refresh_count += 1
        print(f"{Fore.CYAN}[cookie] Opening real Chrome to obtain a fresh cf_clearance "
              f"({self._cookie_refresh_count}/{self.cookie_refresh_max})...")
        return self._refresh_cookies_from_browser()

    def _refresh_cookies_from_browser(self) -> bool:
        """Open the site in a real (visible) Chrome, wait for the Cloudflare
        challenge to pass and the translation page to load (txtOriginal present),
        capture the cookies, and save them. Leaves the driver ON the loaded page
        so callers can continue immediately. Returns True if fresh cookies stored.
        """
        from selenium.webdriver.common.by import By

        self._ensure_driver()
        driver = self.driver
        driver.get(self.site)
        self._inject_cookies(driver)
        driver.get(self.site)

        deadline = time.time() + self.cookie_wait
        loaded = None
        while time.time() < deadline:
            cookies = {c["name"]: c["value"] for c in driver.get_cookies()}
            try:
                # Page fully loaded past the challenge (txtOriginal present).
                driver.find_element(By.ID, "txtOriginal")
                loaded = cookies
                break
            except Exception:
                pass
            time.sleep(1)

        if loaded is None:
            # Page didn't finish loading, but we may still have captured a cookie.
            cookies = {c["name"]: c["value"] for c in driver.get_cookies()}
            if cookies.get("cf_clearance"):
                loaded = cookies
            else:
                print(f"{Fore.RED}[cookie] No page load / cf_clearance within {self.cookie_wait}s.")
                return False
        self._save_cookies(loaded)
        print(f"{Fore.GREEN}[cookie] Fresh cf_clearance captured and saved.")
        return True

    def _web_call(self, text: str, timeout: int = 120) -> str:
        from selenium.webdriver.common.by import By
        from selenium.webdriver.support.ui import WebDriverWait
        from selenium.webdriver.support import expected_conditions as EC

        self._ensure_driver()
        driver = self.driver
        driver.get(self.site)

        # Inject cookies after first navigation (page context exists), then reload
        self._inject_cookies(driver)
        driver.get(self.site)

        wait = WebDriverWait(driver, 30)
        try:
            input_el = wait.until(EC.presence_of_element_located((By.ID, "txtOriginal")))
        except Exception:
            # Cloudflare challenge present — try to auto-refresh the cookie once
            # (the user may need to click the challenge). On success the driver is
            # ALREADY on the loaded page, so we continue from there — no reload.
            print(f"{Fore.YELLOW}[web] Cloudflare challenge detected. Attempting cookie refresh...")
            if not self._maybe_refresh_cookies():
                raise CloudflareBlockedError(
                    "Translation site blocked (Cloudflare challenge) and cookie "
                    "refresh limit reached. Update translate_cookies under Settings "
                    "or click 'Continue' on the book."
                )
            try:
                input_el = wait.until(EC.presence_of_element_located((By.ID, "txtOriginal")))
            except Exception:
                raise CloudflareBlockedError(
                    "Translation site still blocked (Cloudflare challenge) after "
                    "cookie refresh. Update translate_cookies under Settings or "
                    "click 'Continue' on the book."
                )
        input_el.clear()
        input_el.send_keys(text)

        # Click the translate button (anchor with btnTranslateClick handler)
        try:
            btn = driver.find_element(By.CSS_SELECTOR, "a[onclick='btnTranslateClick()'], a[onclick=\"btnTranslateClick()\"]")
        except Exception:
            btn = driver.find_element(By.CSS_SELECTOR, "a.btn.cyan")

        output_el = wait.until(EC.presence_of_element_located((By.ID, "txtTranslation")))

        # Clicking too fast can result in no translation. Wait a moment before the
        # first click, then retry the click (after a delay) if no result appears.
        max_attempts = max(1, self.max_retries)
        for attempt in range(1, max_attempts + 1):
            if attempt == 1:
                time.sleep(self.click_delay)
            else:
                print(f"{Fore.YELLOW}[web] attempt {attempt}: no result yet, "
                      f"retrying click in {self.retry_delay}s...")
                time.sleep(self.retry_delay)

            try:
                btn.click()
            except Exception as e:
                print(f"{Fore.RED}[web] click failed on attempt {attempt}: {e}")
                continue

            # Poll until the translation text appears
            deadline = time.time() + timeout
            while time.time() < deadline:
                value = output_el.get_attribute("value") or ""
                if value.strip():
                    # Successful load also proves the current cookie works — refresh it
                    try:
                        self._save_cookies({c["name"]: c["value"] for c in driver.get_cookies()})
                    except Exception:
                        pass
                    return value.strip()
                time.sleep(1)

        raise RuntimeError(
            f"Web translation timed out after {max_attempts} click attempts "
            f"({self.click_delay}s pre-click delay, {self.retry_delay}s retry delay, "
            f"{self.max_retries} max retries)."
        )

    def translate(self, text: str, method: str = "api") -> str:
        """Translate a block of text. Falls back api -> web automatically.
        Logs which method actually produced the result (one line per call)."""
        text = (text or "").strip()
        if not text:
            return ""
        chunks = self._chunk_text(text)
        results = []
        used = None
        for i, chunk in enumerate(chunks):
            if method == "api":
                try:
                    results.append(self._api_call(chunk))
                    used = "api"
                except Exception as e:
                    print(f"{Fore.YELLOW}API translate failed, falling back to web: {e}")
                    results.append(self._web_call(chunk))
                    used = "web"
            else:
                results.append(self._web_call(chunk))
                used = "web"
            if i < len(chunks) - 1:
                time.sleep(random.uniform(0.5, 1.5))
        if used:
            print(f"[{used}] translated {len(text)} chars")
        return "\n".join(results)

    def _chunk_text(self, text: str) -> list[str]:
        """Split text into chunks no larger than MAX_CHUNK_CHARS at paragraph boundaries."""
        paragraphs = text.split("\n")
        chunks = []
        current = []
        current_len = 0
        for p in paragraphs:
            if current and current_len + len(p) + 1 > MAX_CHUNK_CHARS:
                chunks.append("\n".join(current))
                current = []
                current_len = 0
            current.append(p)
            current_len += len(p) + 1
        if current:
            chunks.append("\n".join(current))
        return chunks

    def close(self):
        if self.driver is not None:
            try:
                self.driver.quit()
            except Exception:
                pass
            self.driver = None

    # --- Book-level translation ---

    def _probe_method(self, method: str) -> bool:
        """Verify the chosen translation method works before starting a run."""
        try:
            probe = self.translate("测试", method=method)
            return bool(probe and probe.strip())
        except Exception as e:
            print(f"{Fore.RED}Translation method '{method}' not usable: {e}")
            return False

    def _build_docx(self, book: dict, output_docx: Path) -> Document:
        """Load an existing DOCX to resume, or create a fresh one with cover/title/author."""
        if output_docx.exists():
            print(f"{Fore.CYAN}Found existing DOCX. Resuming...")
            return Document(str(output_docx))
        doc = Document()
        for p in list(doc.paragraphs):
            p._element.getparent().remove(p._element)
        cover_path = self._find_cover_image(output_docx.parent, book["id"], book["seo_title_basic"])
        if cover_path:
            img_p = doc.add_paragraph()
            img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = img_p.add_run()
            run.add_picture(str(cover_path), width=Inches(3.5))
        doc.add_heading(book["title"], level=0)
        if book.get("author"):
            p = doc.add_paragraph()
            run_label = p.add_run("Tác giả: ")
            run_label.italic = True
            run_author = p.add_run(book["author"])
            run_author.bold = True
            run_author.italic = True
        return doc

    def _translate_run(self, book_id: int, method: str, chapters_to_process: list,
                       output_docx: Path, task: TranslateTask):
        """Shared loop: translate each chapter, append to DOCX, update statuses.
        Original chapter content is NEVER overwritten — translated text only goes to DOCX."""
        corrections = self.db.get_book_corrections(book_id)
        total = len(chapters_to_process)
        task.set(total=total, message="Probing translation method...")

        if not self._probe_method(method):
            self.db.update_book_status(book_id, "failed")
            task.set(active=False, message=f"Translation method '{method}' failed.")
            return
        task.set(message="Probe OK, starting chapters...")

        doc = self._build_docx(self.db.get_book(book_id), output_docx)
        success_count = 0
        fail_count = 0

        try:
            for idx, ch in enumerate(chapters_to_process, 1):
                if task.cancelled:
                    task.set(message="Translation cancelled")
                    break
                task.set(current_index=idx, current_title=ch["chapter_title"])
                try:
                    raw_title = ch["chapter_title"] or f"Chương {ch['chapter_order']}"
                    raw_content = ch.get("chapter_content") or ""

                    # ONE request per chapter: send title + content together so we
                    # don't make two requests. The first line of the output is the
                    # translated title; the rest is the translated content.
                    combined = f"{raw_title}\n{raw_content}" if raw_content else raw_title
                    translated = self.translate(combined, method=method)
                    if "\n" in translated:
                        translated_title, translated_content = translated.split("\n", 1)
                    else:
                        translated_title, translated_content = translated, ""
                    # Safety: if the content came back empty (split edge case),
                    # translate it separately so the chapter is not lost.
                    if raw_content and not translated_content.strip():
                        print(f"{Fore.YELLOW}[run] Title+content split empty; "
                              f"translating content separately.")
                        translated_content = self.translate(raw_content, method=method)

                    translated_title = apply_corrections(translated_title, corrections)
                    translated_content = apply_corrections(translated_content, corrections)

                    doc.add_heading(translated_title, level=1)
                    for para in translated_content.split("\n"):
                        para = para.strip()
                        if para:
                            doc.add_paragraph(para)

                    self.db.update_chapter_status(ch["id"], "completed", file_path=output_docx.name)
                    success_count += 1
                    task.set(success_count=success_count, message=f"Translated: {translated_title}")
                    print(f"{Fore.GREEN}Translated: {translated_title} ({idx}/{total})")
                except CloudflareBlockedError as e:
                    # Cloudflare needs a new session. Mark the CURRENT chapter
                    # failed and STOP — do NOT move on to the next chapter (it
                    # would fail too). The user refreshes the cookie and clicks
                    # 'Continue', which retries failed chapters into the DOCX.
                    self.db.update_chapter_status(ch["id"], "failed", error_message=str(e))
                    fail_count += 1
                    task.set(fail_count=fail_count,
                             message=f"Cloudflare blocked chapter {ch['chapter_order']} — stopped. "
                                     f"Refresh the cookie and click Continue.")
                    print(f"{Fore.RED}Cloudflare blocked chapter {ch['chapter_order']} — "
                          f"STOPPING run. Refresh the cookie and click Continue.")
                    task.set(blocked=True)
                    break
                except Exception as e:
                    self.db.update_chapter_status(ch["id"], "failed", error_message=str(e))
                    fail_count += 1
                    task.set(fail_count=fail_count, message=f"Failed chapter {ch['chapter_order']}: {e}")
                    print(f"{Fore.RED}Chapter {ch['chapter_order']} failed: {e}")

                if idx % 10 == 0:
                    try:
                        doc.save(output_docx)
                        print(f"{Fore.CYAN}--- Checkpoint: DOCX saved ---")
                    except Exception as e:
                        print(f"{Fore.RED}Checkpoint save failed: {e}")

                time.sleep(random.uniform(0.5, 1.5))
        finally:
            try:
                doc.save(output_docx)
            except Exception as e:
                print(f"{Fore.RED}Final save failed: {e}")

        return success_count, fail_count

    def translate_book(self, book_id: int, method: str = "api", run_async: bool = True,
                       max_chapters: int = None):
        """Translate pending chapters into the main {book_id}_{seo}.docx.
        Mirrors downloader.run(): only processes 'pending' chapters, saves progress,
        and updates book status like a download (completed/paused/cancelled/...).
        """
        book = self.db.get_book(book_id)
        if not book:
            raise ValueError("Book not found")

        if run_async:
            thread = threading.Thread(target=self.translate_book,
                                      args=(book_id, method, False, max_chapters), daemon=True)
            thread.start()
            return {"message": f"Translation started for {book['title']}"}

        task = TranslateTask(book_id)
        task.set(method=method, message=f"Starting translation of {book['title']}")
        register_task(task)
        self.db.update_book_status(book_id, "in_progress", total_chapters=book.get("total_chapters"))

        all_chapters = self.db.get_chapters_by_book(book_id)
        chapters_to_process = [c for c in all_chapters if c["download_status"] == "pending"]

        if not chapters_to_process:
            print(f"{Fore.GREEN}All chapters already translated for '{book['title']}'!")
            self.close()
            unregister_task(book_id)
            task.set(active=False, message="All chapters already translated.")
            return

        completed_before = len(all_chapters) - len(chapters_to_process)
        print(f"📚 {len(all_chapters)} total chapters, {completed_before} done, "
              f"{len(chapters_to_process)} remaining.")

        limited = False
        if max_chapters and len(chapters_to_process) > max_chapters:
            chapters_to_process = chapters_to_process[:max_chapters]
            limited = True
            print(f"🔢 Session limited to {max_chapters} chapters.")

        base_dir = Path(__file__).parent.parent.parent
        save_dir = base_dir / (self.db.get_setting("book_path") or TRUYENWIKI["book_path"])
        os.makedirs(save_dir, exist_ok=True)
        output_docx = save_dir / f"{book_id}_{book['seo_title_basic']}.docx"

        try:
            success, fail = self._translate_run(book_id, method, chapters_to_process, output_docx, task)
        except Exception as e:
            print(f"{Fore.RED}Translation run failed: {e}")
            task.set(message=f"Translation run failed: {e}")
            self.db.update_book_status(book_id, "failed")
            self.close()
            unregister_task(book_id)
            return

        if task.cancelled:
            new_status = "cancelled"
        elif task.progress.get("blocked"):
            # Stopped at the current chapter because Cloudflare required a new
            # session. Not a real completion — behave like paused so the UI shows
            # Continue / Translate All after the cookie is refreshed.
            new_status = "paused"
        elif limited:
            new_status = "paused"
        elif fail == 0:
            new_status = "completed"
        elif success == 0:
            new_status = "failed"
        else:
            new_status = "completed_with_errors"

        total_downloaded = completed_before + success
        total_chapters = len(chapters_to_process)
        self.db.update_book_status(book_id, new_status, downloaded_chapters=total_downloaded)
        task.set(message=f"Finished: {success}/{total_chapters} chapters translated ({new_status})")
        task.set(active=False)
        self.close()
        unregister_task(book_id)

    def retranslate_book(self, book_id: int, method: str = "api", run_async: bool = True):
        """Re-translate failed chapters into a separate {book_id}_{seo}_retranslate.docx.
        Mirrors downloader.run_redownload(). Original content is never overwritten.
        """
        book = self.db.get_book(book_id)
        if not book:
            raise ValueError("Book not found")

        if run_async:
            thread = threading.Thread(target=self.retranslate_book,
                                      args=(book_id, method, False), daemon=True)
            thread.start()
            return {"message": f"Re-translation started for {book['title']}"}

        task = TranslateTask(book_id)
        task.set(method=method, message=f"Re-translating failed chapters of {book['title']}")
        register_task(task)
        self.db.update_book_status(book_id, "in_progress", total_chapters=book.get("total_chapters"))

        all_chapters = self.db.get_chapters_by_book(book_id)
        chapters_to_process = [c for c in all_chapters if c["download_status"] == "failed"]

        if not chapters_to_process:
            print(f"{Fore.GREEN}No failed chapters to re-translate.")
            self.close()
            unregister_task(book_id)
            task.set(active=False, message="No failed chapters to re-translate.")
            return

        base_dir = Path(__file__).parent.parent.parent
        save_dir = base_dir / (self.db.get_setting("book_path") or TRUYENWIKI["book_path"])
        os.makedirs(save_dir, exist_ok=True)
        output_docx = save_dir / f"{book_id}_{book['seo_title_basic']}_retranslate.docx"

        try:
            success, fail = self._translate_run(book_id, method, chapters_to_process, output_docx, task)
        except Exception as e:
            print(f"{Fore.RED}Re-translation run failed: {e}")
            task.set(message=f"Re-translation run failed: {e}")
            self.db.update_book_status(book_id, "failed")
            self.close()
            unregister_task(book_id)
            return

        # Status mirrors downloader.run_redownload()
        fresh = self.db.get_chapters_by_book(book_id)
        remaining_pending = len([c for c in fresh if c["download_status"] == "pending"])
        if task.cancelled:
            new_status = "cancelled"
        elif remaining_pending:
            new_status = "paused"
        elif fail == 0:
            new_status = "completed"
        else:
            new_status = "completed_with_errors"
        downloaded = len([c for c in fresh if c["download_status"] == "completed"])
        self.db.update_book_status(book_id, new_status, downloaded_chapters=downloaded)
        task.set(message=f"Finished re-translation: {success}/{len(chapters_to_process)} ({new_status})")
        task.set(active=False)
        self.close()
        unregister_task(book_id)

    def _find_cover_image(self, save_dir, book_id: int, seo_basic: str) -> str | None:
        basename = f"{book_id}_{seo_basic}"
        for ext in [".jpg", ".jpeg", ".png", ".gif", ".webp"]:
            path = save_dir / f"{basename}{ext}"
            if path.exists():
                return str(path)
        return None


def translate_book(book_id: int, method: str = "api", run_async: bool = True,
                   max_chapters: int = None):
    """Module-level helper so the API router doesn't need to instantiate."""
    return Translator().translate_book(book_id, method=method, run_async=run_async,
                                       max_chapters=max_chapters)


def retranslate_book(book_id: int, method: str = "api", run_async: bool = True):
    """Module-level helper for re-translating failed chapters."""
    return Translator().retranslate_book(book_id, method=method, run_async=run_async)


def cancel_all_translations():
    with _lock:
        ids = list(_active_translations.keys())
    for bid in ids:
        cancel_translation(bid)


def continue_book(book_id: int, method: str = "api", run_async: bool = True):
    """Reset failed chapters back to 'pending' and resume them into the MAIN DOCX.
    Use this after updating a stale cf_clearance cookie: it appends the retried
    chapters to {book_id}_{seo}.docx instead of a separate _retranslate.docx."""
    t = Translator()
    try:
        reset = t.db.reset_failed_chapters(book_id)
        if reset == 0:
            return {"message": "No failed chapters to continue.", "reset": 0}
        t.close()
        return translate_book(book_id, method=method, run_async=run_async)
    finally:
        t.close()


def refresh_cookies_manually() -> dict:
    """Open a real Chrome and capture a fresh cf_clearance cookie. Returns result info."""
    t = Translator()
    try:
        ok = t._refresh_cookies_from_browser()
        return {"ok": ok, "has_cookies": bool(t.cookies)}
    finally:
        t.close()
